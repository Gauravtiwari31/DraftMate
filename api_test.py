"""Offline API regression test: api_test.py — provider calls are faked, no API keys or quota used."""
import io
import json
import os
import sys
import tempfile

TMP = tempfile.mkdtemp()
os.environ["CACHE_DIR"] = TMP  # keep demo_cache untouched

import llm as llm_mod

CALLS = []


def fake_call(tier, system, user, max_tokens):
    CALLS.append((tier, max_tokens))
    if "route Indian legal" in system:
        return "general"
    if "analyse the writing style" in system:
        return json.dumps({"tone": "formal", "heading_pattern": "numbered", "defined_terms_style": "caps",
                           "clause_patterns": ["a"], "formatting_rules": ["b"]})
    if "Extract every legal citation" in system:
        return "[]"
    if "convert a spoken" in system:
        return json.dumps({"action": "delete", "needs_confirm": False})
    return "# Draft\n\n1. **Bold** clause text here for the test."


llm_mod._call = fake_call

from fastapi.testclient import TestClient
from docx import Document
import yaml
import main
import verify
import export_docx

c = TestClient(main.app, raise_server_exceptions=False)
results = []


def check(name, cond, detail=""):
    results.append(cond)
    print(f"{'PASS' if cond else 'FAIL'} {name} {detail}")


def last_call():
    return CALLS[-1] if CALLS else (None, None)


def main_():
    uploads_before = set(os.listdir(main.docs_mod.UPLOAD_DIR))

    r = c.get("/health")
    check("health", r.status_code == 200)
    r = c.get("/skills")
    check("skills 19 total / 9 real", len(r.json()) == 19 and sum(not s["stub"] for s in r.json()) == 9)
    ids = {s["id"] for s in r.json() if not s["stub"]}
    check("case_analysis, due_diligence, petition_drafting ready", {"case_analysis", "due_diligence", "petition_drafting"} <= ids)
    r = c.get("/models")
    m = {x["tier"]: x for x in r.json()}
    check("/models lists 3 tiers with labels", r.status_code == 200 and set(m) == {"quick", "balanced", "advanced"}
          and all(x["label"] and x["label"] != "Not configured" for x in m.values()), str([x["label"] for x in m.values()]))
    r = c.post("/improve", json={"message": "@legal_notice tenant not paying"})
    check("/improve returns prompt keeping @tag", r.status_code == 200 and r.json()["message"].startswith("@legal_notice"), r.text[:80])
    r = c.post("/improve", json={"message": " "})
    check("/improve empty -> 422", r.status_code == 422)
    r = c.post("/chat", json={"message": "@legal_notice rent due", "doc_ids": []})
    check("chat @legal_notice", r.status_code == 200 and r.json()["skill"] == "legal_notice")
    with open("samples/reference_nda.docx", "rb") as f:
        r = c.post("/upload", files={"file": ("reference_nda.docx", f)}, data={"role": "reference"})
    check("upload reference docx + style", r.status_code == 200 and "style" in r.json(), str(r.status_code))
    with open("samples/scanned.pdf", "rb") as f:
        r = c.post("/upload", files={"file": ("scanned.pdf", f)}, data={"role": "source"})
    check("scanned pdf -> 422", r.status_code == 422, str(r.status_code))
    r = c.post("/voice", json={"transcript": "delete the draft", "draft": "x"})
    check("voice delete forces confirm", r.status_code == 200 and r.json()["needs_confirm"] is True)
    r = c.post("/verify", json={"text": "no citations"})
    check("verify no citations -> []", r.status_code == 200 and r.json() == [])

    # regressions
    r = c.post("/chat", json={"message": "@rent_agreement draft a rent agreement", "doc_ids": []})
    check("stub skill via @tag -> 422", r.status_code == 422, str(r.status_code))
    r = c.post("/chat", json={"message": "draft", "skill": "rent_agreement"})
    check("stub skill via payload -> 422", r.status_code == 422, str(r.status_code))
    r = c.post("/chat", json={"message": "draft a notice", "skill": "leg"})
    check("unknown skill id ignored", r.status_code == 200 and r.json()["skill"] != "leg", str(r.json().get("skill")))
    r = c.post("/voice", json={"transcript": "undo", "draft": None})
    check("voice draft=null ok", r.status_code == 200, str(r.status_code))
    r = c.post("/chat", json={"message": 123})
    check("non-string message -> 422", r.status_code == 422, str(r.status_code))
    CALLS.clear()
    c.post("/chat", json={"message": "@legal_notice x", "tier": "quick"})
    check("explicit model tier honoured", last_call()[0] == "quick", str(last_call()))
    CALLS.clear()
    c.post("/chat", json={"message": "@summarise y", "tier": "auto"})
    check("auto tier uses skill default", last_call()[0] == "balanced", str(last_call()))
    CALLS.clear()
    c.post("/edit", json={"instruction": "shorten", "draft": "word " * 8000})  # ~10k tokens
    check("edit budget fits long draft", (last_call()[1] or 0) >= 40000 // 4, str(last_call()))
    r = c.post("/edit", json={"instruction": "shorten", "draft": "word " * 20000})
    check("edit of oversize draft -> 413 (not silent truncation)", r.status_code == 413, str(r.status_code))
    check("lookalike domain not Verified", verify.status("full", "https://indiankanoon.org.evil.example/x") != "Verified")
    check("real domain Verified", verify.status("full", "https://indiankanoon.org/doc/1") == "Verified")
    r = c.post("/upload", files={"file": ("x.html", io.BytesIO(b"<p>" + b"legal text " * 10 + b"</p>"))}, data={"role": "source"})
    check("non pdf/docx/txt/md upload -> 415", r.status_code == 415, str(r.status_code))
    r = c.post("/upload", files={"file": ("big.txt", io.BytesIO(b"a" * (5 * 1024 * 1024 + 1)))}, data={"role": "source"})
    check("upload > 5 MB -> 413", r.status_code == 413, str(r.status_code))
    r = c.post("/export", json={"draft": "# T\n1. **Bold** term", "format": "docx"})
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    check("export docx, bold rendered", r.status_code == 200 and "1. Bold term" in texts, str(texts))
    check("no files left in uploads/", set(os.listdir(main.docs_mod.UPLOAD_DIR)) == uploads_before)
    # corrupt every cache entry: the next call must refetch, not 500
    for e in os.scandir(TMP):
        open(e.path, "w").write("{not json")
    r = c.post("/chat", json={"message": "@legal_notice rent due", "doc_ids": []})
    check("corrupt cache entry refetched", r.status_code == 200, str(r.status_code))
    p = io.BytesIO()
    export_docx.export_docx("## **Head**", p)
    check("heading strips **", Document(p).paragraphs[0].text == "Head")
    doc = main.docs_mod.ingest(main.Path("samples/big_50page.docx"), "source")
    picked = main.docs_mod.retrieve("payment", [doc["doc_id"]], k=4)
    check("retrieved chunks keep document order", [doc["chunks"].index(x) for x in picked] == sorted(doc["chunks"].index(x) for x in picked))
    with open("samples/reference_nda.docx", "rb") as f:
        nda = c.post("/upload", files={"file": ("reference_nda.docx", f)}, data={"role": "source"}).json()["doc_id"]
    check("summarise sends every clause of a small doc",
          main.docs_mod.retrieve("summarise", [nda], whole=True) == main.docs_mod.get(nda)["chunks"])
    try:
        envs =[e["key"] for e in yaml.safe_load(open("render.yaml"))["services"][0]["envVars"]]
        check("render.yaml parses with ALLOWED_ORIGIN", "ALLOWED_ORIGIN" in envs)
    except Exception as e:
        check("render.yaml parses", False, f"{type(e).__name__}")

    print(f"\n{sum(results)}/{len(results)} passed")
    print("RESULT:", "PASS" if all(results) else "FAIL")
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main_())
