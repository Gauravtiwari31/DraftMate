import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_runs(paragraph, text):
    # **bold** spans become bold runs instead of literal asterisks
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = i % 2 == 1


def export_docx(draft, path):
    """Write the draft as DOCX to a path or a binary stream."""
    doc = Document()
    for line in draft.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 4))
        elif re.match(r"^\d+(\.\d+)*[\.\)]\s+\S", stripped):
            p = doc.add_paragraph()
            _add_runs(p, stripped)
            p.paragraph_format.left_indent = Pt(18)
        else:
            _add_runs(doc.add_paragraph(), stripped)
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = "AI-assisted draft, requires professional review"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs:
        r.font.size = Pt(8)
    doc.save(path)
    return path
