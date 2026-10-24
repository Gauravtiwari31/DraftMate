"""Generate the sample documents in samples/ (all fictional or public)."""
from pathlib import Path

SAMPLES = Path(__file__).parent.parent / "samples"
SAMPLES.mkdir(exist_ok=True)


def make_pdf(lines, out: Path):
    content = "BT /F1 12 Tf 72 720 Td\n"
    for line in lines:
        safe = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({safe}) Tj 0 -16 Td\n"
    content += "ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    pdf = "%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(pdf.encode()))
        pdf += f"{i} 0 obj\n{body}\nendobj\n"
    xref_pos = len(pdf.encode())
    pdf += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
    out.write_bytes(pdf.encode("latin-1"))


def make_nda(out: Path):
    from docx import Document
    doc = Document()
    doc.add_heading("MUTUAL NON-DISCLOSURE AGREEMENT", 0)
    doc.add_heading("1. Purpose", level=2)
    doc.add_paragraph(
        "This Agreement is made on 1 April 2025 between Alpha Analytics Pvt. Ltd., having its "
        "registered office at Mumbai, and Beta Systems Pvt. Ltd., to evaluate a potential business "
        "relationship concerning a data-processing platform."
    )
    doc.add_heading("2. Definition of Confidential Information", level=2)
    doc.add_paragraph(
        "Confidential Information means any non-public technical, business or financial information "
        "disclosed by either party, whether orally, in writing or in electronic form, that is "
        "designated as confidential or that reasonably ought to be understood as confidential."
    )
    doc.add_heading("3. Obligations of the Receiving Party", level=2)
    doc.add_paragraph(
        "The Receiving Party shall (a) use the Confidential Information solely for the Purpose; "
        "(b) protect it with at least the same care it applies to its own confidential information; "
        "(c) not disclose it to any third party without prior written consent."
    )
    doc.add_heading("4. Term and Termination", level=2)
    doc.add_paragraph(
        "This Agreement shall remain in force for two (2) years from the Effective Date and may be "
        "terminated by either party on thirty (30) days written notice."
    )
    doc.add_heading("5. Governing Law", level=2)
    doc.add_paragraph(
        "This Agreement shall be governed by the laws of India and the courts at Mumbai shall have "
        "exclusive jurisdiction. Nothing herein shall be construed as waiving either party's "
        "obligations under Section 27 of the Indian Contract Act, 1872."
    )
    doc.save(out)


def make_source_letter(out: Path):
    from docx import Document
    doc = Document()
    doc.add_heading("INSTRUCTION LETTER — RENT ARREARS MATTER (FICTIONAL)", 0)
    doc.add_paragraph(
        "Dear Counsel,\n\nOur client, Mr. R. Deshmukh, owns a flat at Kothrud, Pune, let out to "
        "Mr. S. Kulkarni since 1 April 2023 at a monthly rent of Rs 15,000 under a registered leave "
        "and licence agreement. No rent has been paid for the months of May through October 2025, "
        "totalling Rs 90,000. Two oral demands and one WhatsApp reminder have been ignored. "
        "The agreement is governed by Maharashtra law; the security deposit of Rs 45,000 is with us. "
        "Please draft a legal notice demanding payment within 15 days, failing which we will pursue "
        "remedies under the Maharashtra Rent Control Act, 1999 and, if required, civil recovery.\n\n"
        "Yours faithfully,\nAlpha Chambers (fictional, for demo use only)"
    )
    doc.save(out)


def make_demo_draft(out: Path):
    from docx import Document
    doc = Document()
    doc.add_heading("DRAFT LEGAL NOTICE (SAMPLE WITH CITATIONS)", 0)
    doc.add_paragraph(
        "1. You are hereby called upon to pay the arrears of rent within fifteen (15) days. "
        "The rights of my client are affected by the law laid down in Kesavananda Bharati v. State "
        "of Kerala, (1973) 4 SCC 225, and the principle of natural justice recognised in Maneka "
        "Gandhi v. Union of India, (1978) 1 SCC 248."
    )
    doc.add_paragraph(
        "2. Reliance is also placed on Balakrishnan Menon v. Corporation of Trivandrum, (1989) 3 "
        "SCC 412, which supports the claim for mesne profits."
    )
    doc.add_paragraph(
        "Note for testers: citation 1 and 2 are real and should verify green; citation 3 (clause 2) "
        "is planted and does not exist — it should show red/Unverified."
    )
    doc.save(out)


def make_big_doc(out: Path):
    from docx import Document
    doc = Document()
    for page in range(50):
        doc.add_heading(f"Chapter {page + 1}: Confidential Data Processing Schedule", level=1)
        for clause in range(4):
            doc.add_paragraph(
                f"{page + 1}.{clause + 1} The parties agree that all records processed under this "
                "schedule shall remain the property of the Disclosing Party and shall be returned "
                "or destroyed upon termination. The Receiving Party shall maintain an audit log "
                "of every access and shall report any breach within seventy-two hours."
            )
    doc.save(out)


if __name__ == "__main__":
    make_nda(SAMPLES / "reference_nda.docx")
    make_pdf([
        "MUTUAL NON-DISCLOSURE AGREEMENT",
        "1. Purpose: evaluation of a joint data-processing venture.",
        "2. Confidential Information covers all non-public technical data.",
        "3. The Receiving Party shall not disclose any Confidential Information.",
        "4. Term: two years from the Effective Date.",
        "5. Governing law: India; jurisdiction: courts at Mumbai.",
    ], SAMPLES / "reference_nda.pdf")
    make_source_letter(SAMPLES / "source_letter.docx")
    make_demo_draft(SAMPLES / "demo_draft.docx")
    make_big_doc(SAMPLES / "big_50page.docx")
    make_pdf([" "], SAMPLES / "scanned.pdf")
    print("samples written to", SAMPLES)
